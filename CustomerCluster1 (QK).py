from SalesDatabaseProject.dal.customer_dao import CustomerDAO
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
from sklearn.cluster import KMeans
import plotly.express as px

df = CustomerDAO.getCustomersForClusering()
print(df)
print(df.columns)


def showHistogram(df, columns):
    plt.figure(1, figsize=(8, 10))
    n = 0
    for column in columns:
        n += 1
        plt.subplot(3, 1, n)
        plt.subplots_adjust(hspace=0.4, wspace=0.4)

        sns.histplot(df[column], bins=32, kde=True)

        plt.title(f'Histogram for {column}')
    plt.show()


showHistogram(df, df.columns[1:])
def elbowMethod(df, columnsForElbow, maxK=10, tolerance=0.9):
    X = df.loc[:, columnsForElbow].values
    kRange = np.arange(1, maxK + 1)
    inertia = []
    for n in kRange:
        model = KMeans(n_clusters=n,
                       init='k-means++',
                       max_iter=500,
                       random_state=42)
        model.fit(X)
        inertia.append(model.inertia_)
    inertia = np.array(inertia)

    # --- Tự động tìm số cụm tối ưu (thuật toán Kneedle) ---
    # 1) Chuẩn hóa K và inertia về đoạn [0, 1] để 2 trục cùng thang đo
    xNorm = (kRange - kRange.min()) / (kRange.max() - kRange.min())
    yNorm = (inertia - inertia.min()) / (inertia.max() - inertia.min() + 1e-12)
    # 2) Đường cong elbow đi xuống, đường thẳng nối điểm đầu - điểm cuối là y = 1 - x
    #    Khoảng cách (theo chiều dọc) từ mỗi điểm tới đường thẳng đó: d = (1 - x) - y
    #    Điểm có d lớn nhất là điểm "gập" rõ nhất
    distance = (1 - xNorm) - yNorm
    maxDistance = distance.max()
    # 3) Khoảng elbow: các K có độ "gập" >= tolerance * độ gập lớn nhất
    #    (mặc định 90%) => các K này gần như tốt ngang nhau, khó lựa chọn
    elbowZone = kRange[distance >= tolerance * maxDistance]
    # 4) Nếu có nhiều K trong khoảng elbow thì ưu tiên K lớn hơn
    optimalK = int(elbowZone.max())

    print("Inertia theo K:", dict(zip(kRange.tolist(), np.round(inertia, 2).tolist())))
    print("Khoảng elbow:", elbowZone.tolist(), "=> chọn K =", optimalK)

    # --- Vẽ biểu đồ ---
    plt.figure(1, figsize=(15, 6))
    plt.plot(kRange, inertia, 'o')
    plt.plot(kRange, inertia, '-', alpha=0.5)
    # Tô vùng khoảng elbow (nếu có nhiều hơn 1 K)
    if len(elbowZone) > 1:
        plt.axvspan(elbowZone.min() - 0.3, elbowZone.max() + 0.3,
                    color='orange', alpha=0.15,
                    label=f'Khoảng elbow: K = {elbowZone.min()} - {elbowZone.max()}')
    # Đường kẻ dọc tại cụm tối ưu
    plt.axvline(x=optimalK, color='red', linestyle='--', linewidth=1.5,
                label=f'Số cụm tối ưu K = {optimalK}')
    plt.plot(optimalK, inertia[optimalK - 1], 'o', color='red', markersize=12)
    plt.annotate(f'K = {optimalK}',
                 xy=(optimalK, inertia[optimalK - 1]),
                 xytext=(optimalK + 0.4, inertia[optimalK - 1] + (inertia.max() - inertia.min()) * 0.1),
                 color='red', fontsize=12, fontweight='bold',
                 arrowprops=dict(arrowstyle='->', color='red'))
    plt.xticks(kRange)
    plt.xlabel('Number of Clusters'), plt.ylabel('Cluster sum of squared distances')
    plt.title('Elbow Method')
    plt.legend()
    plt.show()

    return optimalK
columns=['Age','Spending_Score']
k = elbowMethod(df, columns)
print("bạn nên gom cụm thành =",k,"cụm")
def runKMeans(X, cluster):
    model = KMeans(n_clusters=cluster,
                   init='k-means++',
                   max_iter=500,
                   random_state=42)
    model.fit(X)
    labels = model.labels_
    centroids = model.cluster_centers_
    y_kmeans = model.fit_predict(X)
    return y_kmeans, centroids, labels


def visualizeKMeans(X, y_kmeans, cluster, title, xlabel, ylabel, colors):
    plt.figure(figsize=(10, 10))
    for i in range(cluster):
        plt.scatter(X[y_kmeans == i, 0],
                    X[y_kmeans == i, 1],
                    s=100,
                    c=colors[i],
                    label='Cluster %i' % (i + 1))
    plt.title(title)
    plt.xlabel(xlabel)
    plt.ylabel(ylabel)
    plt.legend()
    plt.show()


# Thực thi thuật toán và trực quan hóa
columns = ['Age', 'Spending_Score']
X = df.loc[:, columns].values
cluster = 4
colors = ["red", "green", "blue", "purple", "black", "pink", "orange"]

y_kmeans, centroids, labels = runKMeans(X, cluster)
print(y_kmeans)
print(centroids)
print(labels)

df["cluster"] = labels

visualizeKMeans(X,
                y_kmeans,
                cluster,
                "Clusters of Customers - Age X Spending Score",
                "Age",
                "Spending Score",
                colors)
def export_df_to_excel(df, filename):
    from openpyxl.styles import Font, PatternFill, Alignment

    # Sắp xếp các cụm theo thứ tự tăng dần để dễ quan sát
    clusters = sorted(df["cluster"].unique())
    total = len(df)

    # --- Sheet thống kê phân bố dữ liệu theo từng cụm ---
    # (bỏ qua cột cluster và các cột mã ID vì thống kê chúng không có ý nghĩa)
    numericCols = [c for c in df.select_dtypes(include="number").columns
                   if c != "cluster" and not c.lower().endswith("id")]
    rows = []
    for c in clusters:
        dfCluster = df[df["cluster"] == c]
        row = {
            "Cụm": f"Cluster {c + 1}",
            "Giá trị cột cluster": c,
            "Số lượng": len(dfCluster),
            "Tỷ lệ (%)": round(len(dfCluster) / total * 100, 2),
        }
        # Trung bình / nhỏ nhất / lớn nhất của các cột số trong từng cụm
        for col in numericCols:
            row[f"{col} (TB)"] = round(dfCluster[col].mean(), 2)
            row[f"{col} (Min)"] = dfCluster[col].min()
            row[f"{col} (Max)"] = dfCluster[col].max()
        rows.append(row)
    # Dòng tổng cộng
    totalRow = {"Cụm": "Tổng cộng", "Giá trị cột cluster": "",
                "Số lượng": total, "Tỷ lệ (%)": 100.0}
    for col in numericCols:
        totalRow[f"{col} (TB)"] = round(df[col].mean(), 2)
        totalRow[f"{col} (Min)"] = df[col].min()
        totalRow[f"{col} (Max)"] = df[col].max()
    rows.append(totalRow)
    dfStats = pd.DataFrame(rows)

    headerFont = Font(bold=True, color="FFFFFF")
    headerFill = PatternFill("solid", fgColor="305496")

    def formatSheet(ws):
        # Tô đậm tiêu đề, cố định dòng tiêu đề, tự giãn độ rộng cột
        for cell in ws[1]:
            cell.font = headerFont
            cell.fill = headerFill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for colCells in ws.columns:
            maxLen = max(len(str(cell.value)) if cell.value is not None else 0 for cell in colCells)
            ws.column_dimensions[colCells[0].column_letter].width = min(maxLen + 3, 40)

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        # Mỗi cụm 1 sheet, xuất theo thứ tự cụm
        for c in clusters:
            sheetName = f"Cluster {c + 1}"
            dfCluster = df[df["cluster"] == c]
            dfCluster.to_excel(writer, sheet_name=sheetName, index=False)
            formatSheet(writer.sheets[sheetName])
            print(f"Đã xuất {sheetName}: {len(dfCluster)} dòng")

        # Sheet cuối cùng: thống kê
        dfStats.to_excel(writer, sheet_name="Thống kê", index=False)
        wsStats = writer.sheets["Thống kê"]
        formatSheet(wsStats)
        for cell in wsStats[wsStats.max_row]:   # tô đậm dòng tổng cộng
            cell.font = Font(bold=True)

    print(f"Đã xuất dữ liệu ra file: {filename}")
    print(dfStats)

export_file= "../results/CustomerClusters.xlsx"
export_df_to_excel(df, export_file)

def export_df_to_html(df, filename):
    import os

    dfHtml = df.copy()
    dfHtml["cluster"] = dfHtml["cluster"].astype(int)
    dfHtml["Cụm"] = dfHtml["cluster"].apply(lambda c: f"Cluster {c + 1}")

    # --- Chia nhóm tuổi: <=35 Thanh niên, <=45 Trung niên, còn lại Cao tuổi ---
    def ageGroup(age):
        if age <= 35:
            return "Thanh niên (≤35)"
        elif age <= 45:
            return "Trung niên (36-45)"
        return "Cao tuổi (>45)"
    ageOrder = ["Thanh niên (≤35)", "Trung niên (36-45)", "Cao tuổi (>45)"]
    ageColors = {ageOrder[0]: "#2a9d8f", ageOrder[1]: "#e9c46a", ageOrder[2]: "#e76f51"}
    dfHtml["Nhóm tuổi"] = dfHtml["Age"].apply(ageGroup)

    clusters = sorted(dfHtml["cluster"].unique())
    clusterNames = [f"Cluster {c + 1}" for c in clusters]
    clusterColors = ["#4361ee", "#f72585", "#4cc9f0", "#7209b7",
                     "#fb8500", "#2a9d8f", "#e63946", "#8d99ae"]
    total = len(dfHtml)
    incomeCol = "Annual_Income"
    scoreCol = "Spending_Score"

    # Biến đếm để chỉ nhúng thư viện plotly.js 1 lần (file html xem được cả khi offline)
    state = {"first": True}

    def figToHtml(fig):
        fig.update_layout(autosize=True, margin=dict(l=50, r=20, t=60, b=50))
        out = fig.to_html(full_html=False, include_plotlyjs=state["first"],
                          default_width="100%", default_height="420px",
                          config={"displaylogo": False, "responsive": True})
        state["first"] = False
        return out

    def tableHtml(data):
        return data.to_html(index=False, classes="data-table", border=0,
                            float_format=lambda x: f"{x:,.2f}")

    # --- 1) Biểu đồ tỷ lệ khách hàng theo cụm ---
    dfCount = dfHtml.groupby("Cụm").size().reindex(clusterNames).reset_index(name="Số khách hàng")
    figPie = px.pie(dfCount, names="Cụm", values="Số khách hàng", hole=0.45,
                    title="Tỷ lệ khách hàng theo cụm",
                    color="Cụm", color_discrete_sequence=clusterColors,
                    category_orders={"Cụm": clusterNames})
    figPie.update_traces(textinfo="percent+value", sort=False)

    # --- 2) Bảng tóm tắt từng cụm ---
    dfSummary = dfHtml.groupby("Cụm").agg(
        **{"Số khách hàng": ("cluster", "size"),
           "Tuổi TB": ("Age", "mean"),
           "Thu nhập TB": (incomeCol, "mean"),
           "Spending Score TB": (scoreCol, "mean")}).reindex(clusterNames).reset_index()
    dfSummary.insert(2, "Tỷ lệ (%)", dfSummary["Số khách hàng"] / total * 100)
    for g in ageOrder:
        dfSummary[g] = dfHtml[dfHtml["Nhóm tuổi"] == g].groupby("Cụm").size() \
            .reindex(clusterNames).fillna(0).astype(int).values

    # Biểu đồ nhóm tuổi trong từng cụm (cột chồng)
    dfAgeCount = dfHtml.groupby(["Cụm", "Nhóm tuổi"]).size().reset_index(name="Số khách hàng")
    figAge = px.bar(dfAgeCount, x="Cụm", y="Số khách hàng", color="Nhóm tuổi",
                    title="Cơ cấu nhóm tuổi trong từng cụm", text="Số khách hàng",
                    category_orders={"Cụm": clusterNames, "Nhóm tuổi": ageOrder},
                    color_discrete_map=ageColors)

    overviewCharts = f'<div class="chart">{figToHtml(figPie)}</div><div class="chart">{figToHtml(figAge)}</div>'

    # --- 3) Mỗi cụm: Độ tuổi vs Income, Độ tuổi vs Spending Score ---
    clusterSections = ""
    for c, name, color in zip(clusters, clusterNames, clusterColors * 10):
        dfC = dfHtml[dfHtml["cluster"] == c]
        charts = ""
        for col, label in [(incomeCol, "Thu nhập (Annual Income)"),
                           (scoreCol, "Spending Score")]:
            fig = px.box(dfC, x="Nhóm tuổi", y=col, color="Nhóm tuổi",
                         points="all", hover_data=["CustomerId", "Age"] if "CustomerId" in dfC else ["Age"],
                         title=f"{name}: Nhóm tuổi vs {label}",
                         category_orders={"Nhóm tuổi": ageOrder},
                         color_discrete_map=ageColors,
                         labels={col: label})
            fig.update_layout(showlegend=False)
            charts += f'<div class="chart">{figToHtml(fig)}</div>'

        ageCounts = dfC["Nhóm tuổi"].value_counts().reindex(ageOrder).fillna(0).astype(int)
        ageText = " · ".join(f"{g}: <b>{n}</b>" for g, n in ageCounts.items())
        clusterSections += f"""
        <section class="cluster" style="border-top-color:{color}">
          <h2><span class="dot" style="background:{color}"></span>{name}
              <small>{len(dfC)} khách hàng ({len(dfC) / total * 100:.1f}%)</small></h2>
          <p class="meta">Tuổi TB <b>{dfC['Age'].mean():.1f}</b> ·
             Thu nhập TB <b>{dfC[incomeCol].mean():,.2f}</b> ·
             Spending Score TB <b>{dfC[scoreCol].mean():.1f}</b><br>{ageText}</p>
          <div class="grid">{charts}</div>
          <details><summary>Xem danh sách khách hàng của {name}</summary>
            {tableHtml(dfC.drop(columns=["cluster"]))}</details>
        </section>"""

    # --- 4) Ghép thành trang HTML ---
    page = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Phân cụm khách hàng</title>
<style>
  body {{ font-family: "Segoe UI", Arial, sans-serif; margin: 0; background: #f4f6fb; color: #222; }}
  header {{ background: #1d3557; color: #fff; padding: 28px 40px; }}
  header h1 {{ margin: 0 0 6px; }}
  main {{ max-width: 1200px; margin: 0 auto; padding: 24px; }}
  .cards {{ display: flex; gap: 16px; flex-wrap: wrap; margin-bottom: 20px; }}
  .card {{ background: #fff; border-radius: 10px; padding: 16px 22px; flex: 1; min-width: 180px;
          box-shadow: 0 1px 4px rgba(0,0,0,.08); }}
  .card b {{ display: block; font-size: 28px; color: #1d3557; }}
  section {{ background: #fff; border-radius: 10px; padding: 20px; margin-bottom: 24px;
            box-shadow: 0 1px 4px rgba(0,0,0,.08); }}
  section.cluster {{ border-top: 6px solid; }}
  h2 small {{ font-weight: normal; color: #666; font-size: 15px; margin-left: 8px; }}
  .dot {{ display: inline-block; width: 14px; height: 14px; border-radius: 50%; margin-right: 8px; }}
  .meta {{ color: #444; line-height: 1.7; }}
  .grid {{ display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 16px; }}
  .chart {{ min-width: 0; }}
  .data-table {{ border-collapse: collapse; width: 100%; font-size: 14px; margin-top: 10px; }}
  .data-table th {{ background: #1d3557; color: #fff; padding: 8px; position: sticky; top: 0; }}
  .data-table td {{ padding: 6px 8px; border-bottom: 1px solid #e5e5e5; text-align: center; }}
  .data-table tr:nth-child(even) td {{ background: #f8f9fc; }}
  details {{ margin-top: 12px; }}
  details > div, details table {{ max-height: 400px; overflow: auto; display: block; }}
  summary {{ cursor: pointer; color: #1d3557; font-weight: 600; }}
  @media (max-width: 900px) {{ .grid {{ grid-template-columns: 1fr; }} main {{ padding: 12px; }} }}
</style>
</head>
<body>
<header>
  <h1>Phân cụm khách hàng (K-Means)</h1>
  <div>Nhóm tuổi: Thanh niên ≤ 35 · Trung niên 36-45 · Cao tuổi &gt; 45</div>
</header>
<main>
  <div class="cards">
    <div class="card">Tổng khách hàng<b>{total}</b></div>
    <div class="card">Số cụm<b>{len(clusters)}</b></div>
    <div class="card">Tuổi trung bình<b>{dfHtml['Age'].mean():.1f}</b></div>
    <div class="card">Spending Score TB<b>{dfHtml[scoreCol].mean():.1f}</b></div>
  </div>
  <section>
    <h2>Tổng quan các cụm</h2>
    <div class="grid">{overviewCharts}</div>
    {tableHtml(dfSummary)}
  </section>
  {clusterSections}
  <section>
    <h2>Toàn bộ dữ liệu phân cụm</h2>
    <details><summary>Xem tất cả {total} khách hàng</summary>
      {tableHtml(dfHtml.sort_values(["cluster"]).drop(columns=["cluster"]))}</details>
  </section>
</main>
<script>
  // Vẽ lại các biểu đồ theo đúng kích thước khung sau khi trang tải xong
  window.addEventListener("load", function () {{
    document.querySelectorAll(".js-plotly-plot").forEach(function (d) {{ Plotly.Plots.resize(d); }});
  }});
</script>
</body>
</html>"""

    folder = os.path.dirname(filename)
    if folder:
        os.makedirs(folder, exist_ok=True)
    with open(filename, "w", encoding="utf-8") as f:
        f.write(page)
    print(f"Đã xuất báo cáo HTML ra file: {os.path.abspath(filename)}")


export_df_to_html(df, "../results/index.html")
